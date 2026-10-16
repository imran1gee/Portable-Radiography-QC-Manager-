from openpyxl.styles import Font, Alignment

SECTIONS = {
    16: "1. Timer and Exposure Switch",
    22: "2. Filtration and Tube Housing",
    27: "3. Technique Factors / Machine Performance",
    37: "4. Safety Factors",
    42: "5. Shielding",
    48: "6. Collimation",
}

CHECKS = {
    18: "a. Deadman type exposure switch",
    19: "b. Preset timer terminates",
    20: "c. Timer accurate",
    21: 'd. Exposure on "zero" setting not possible',
    23: "a. Adequate filtration",
    24: "b. Tube head stability",
    25: "c. Tube head leakage",
    26: "d. Standby leakage from cap discharge unit",
    28: "a. Technique factors indicated",
    29: "b. Kilovolt indicator accurate",
    30: "c. Exposures linear",
    31: "d. Exposures reproducible",
    32: "e. Automatic exposure control",
    33: "i. Exposure reproducible",
    34: "ii. AEC mode indicated",
    35: "iii. Back-up timer functions",
    36: "f. Entrance exposures within specified limits",
    38: "a. Visual and oral communication w/patient",
    39: "b. Beam-on indicator",
    40: "c. Protective apparel provided (as required)",
    41: "d. Technique guides",
    43: "a. Operator shielding adequate",
    44: "b. Exposure switch mounted behind shield",
    45: "c. Mobile equipment - 6 foot exposure cord",
    46: "d. Room shielding adequate",
    47: "e. Gonadal shielding provided (as req.)",
    49: "a. Collimating equipment adequate",
    50: "b. X-ray / light alignment adequate",
    51: "c. Beam centering satisfactory",
    52: "d. SID indicator accurate",
}

def populate_state_form(ws):
    for row, title in SECTIONS.items():
        cell = ws[f"C{row}"]
        cell.value = title
        cell.font = Font(name="Arial", size=9, bold=True)

    for row, label in CHECKS.items():
        cell = ws[f"D{row}"]
        cell.value = label
        cell.font = Font(name="Arial", size=8)
        cell.alignment = Alignment(
            vertical="center", wrap_text=False
        )

    ws["B15"] = "Check if non-compliant"
    ws["I15"] = "mA Reproducibility and Linearity:"
    ws["I26"] = "kVp Accuracy"
    ws["N26"] = "Beam Quality"
    ws["K53"] = "COLLIMATION TYPE"
    ws["J60"] = "ADDITIONAL SYSTEM FEATURES"

    # No compliance marks until QC results are available.


def add_ma_measurement_table(ws):
    """Reproduce the original mA measurement table structure."""
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

    thin = Side(style="thin", color="555555")
    green = PatternFill("solid", fgColor="CCFFCC")

    # Remove only merges inside the measurement table.
    for merged in list(ws.merged_cells.ranges):
        if (
            merged.min_row >= 16
            and merged.max_row <= 24
            and merged.min_col >= 9
            and merged.max_col <= 17
        ):
            ws.unmerge_cells(str(merged))

    for row in range(16, 25):
        for col in range(9, 18):
            cell = ws.cell(row, col)
            cell.border = Border(
                left=thin, right=thin, top=thin, bottom=thin
            )
            cell.alignment = Alignment(
                horizontal="center", vertical="center",
                shrink_to_fit=True
            )

    ws.merge_cells("I16:K16")
    ws.merge_cells("L16:O16")
    ws.merge_cells("P16:P18")
    ws.merge_cells("Q16:Q18")

    ws["I16"] = "Pre-Set"
    ws["L16"] = "Measured mR"
    ws["P16"] = "C.V."
    ws["Q16"] = "mR/mAs Linearity"

    labels = {
        "I18": "mA",
        "J18": "Time (s)",
        "K18": "mAs",
        "L18": "mR1",
        "M18": "mR2",
        "N18": "mR3",
        "O18": "mR4",
    }

    for address, value in labels.items():
        ws[address] = value

    for row in range(19, 23):
        for col in range(9, 16):
            ws.cell(row, col).fill = green

    for row in range(16, 19):
        for col in range(9, 18):
            cell = ws.cell(row, col)
            if cell.__class__.__name__ != "MergedCell":
                cell.font = Font(name="Arial", size=8, bold=True)
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True
                )

    # Preserve the original checklist and leave measurements blank.


def add_kvp_and_beam_quality_tables(ws):
    """Format the kVp accuracy and beam quality sections."""
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

    thin = Side(style="thin", color="555555")
    green = PatternFill("solid", fgColor="CCFFCC")

    # Remove extra mA table borders in rows 23-24.
    for row in (23, 24):
        for col in range(9, 18):
            ws.cell(row, col).border = Border()
            ws.cell(row, col).fill = PatternFill(fill_type=None)

    # Original mA table information.
    ws["I15"] = "mA Reproducibility and Linearity:"
    ws["O15"] = "kVp:"
    ws["P15"] = "SDD:"
    # Actual kVp and SDD values remain blank.

    # kVp accuracy table: I27:L35.
    ws["I26"] = "kVp Accuracy"
    for row in range(27, 36):
        for col in range(9, 13):
            cell = ws.cell(row, col)
            cell.border = Border(
                left=thin, right=thin, top=thin, bottom=thin
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                shrink_to_fit=True
            )

    ws["J27"] = "Selected kVp"
    ws["J29"] = "Measured kVp"
    ws["I27"] = "mAs"
    ws["I34"] = "Avg. kVp"
    ws["I35"] = "P/F"

    for row in range(28, 34):
        for col in range(10, 13):
            ws.cell(row, col).fill = green

    # Beam quality table: N27:Q31.
    ws["N26"] = "Beam Quality"

    for row in range(27, 32):
        for col in range(14, 18):
            cell = ws.cell(row, col)
            cell.border = Border(
                left=thin, right=thin, top=thin, bottom=thin
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                shrink_to_fit=True
            )

    for address, value in {
        "N27": "kVp:",
        "N28": "mAs:",
        "N29": "msec:",
        "P30": "HVL:",
        "P31": "P/F:",
    }.items():
        ws[address] = value

    for row in (27, 28, 29):
        ws[f"O{row}"].fill = green

    # Style headings without inserting test results.
    for address in ("I15", "I26", "N26", "J27", "J29"):
        ws[address].font = Font(name="Arial", size=8, bold=True)

    for address in ("I15", "I26", "N26"):
        ws[address].alignment = Alignment(
            vertical="center", shrink_to_fit=True
        )
