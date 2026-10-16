import json
import warnings
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Alignment, Protection

SOURCE = Path.home() / "Desktop/Portable Radiography QC manager/QC 2026-08-13 HP Comer Rad Portable 4.xlsx"
DEST = Path("services/layouts/state_form.json")

def color_dict(c):
    if c is None:
        return None
    return {
        "type": c.type,
        "rgb": str(c.rgb) if c.type == "rgb" else None,
        "indexed": c.indexed if c.type == "indexed" else None,
        "theme": c.theme if c.type == "theme" else None,
        "tint": c.tint,
        "auto": c.auto if c.type == "auto" else None,
    }

def side_dict(s):
    return {
        "style": s.style,
        "color": color_dict(s.color),
    }

def font_dict(f):
    return {
        "name": f.name,
        "sz": f.sz,
        "b": f.b,
        "i": f.i,
        "u": f.u,
        "strike": f.strike,
        "color": color_dict(f.color),
        "vertAlign": f.vertAlign,
    }

def style_dict(c):
    f = c.font
    a = c.alignment
    b = c.border
    fill = c.fill

    return {
        "font": font_dict(f),
        "alignment": {
            "horizontal": a.horizontal,
            "vertical": a.vertical,
            "textRotation": a.textRotation,
            "wrapText": a.wrapText,
            "shrinkToFit": a.shrinkToFit,
            "indent": a.indent,
        },
        "fill": {
            "patternType": fill.patternType,
            "fgColor": color_dict(fill.fgColor),
            "bgColor": color_dict(fill.bgColor),
        },
        "border": {
            side: side_dict(getattr(b, side))
            for side in ("left", "right", "top", "bottom")
        },
        "number_format": c.number_format,
    }

with warnings.catch_warnings():
    warnings.filterwarnings("ignore", message="wmf image format")
    wb = load_workbook(SOURCE)

ws = wb["State Form"]

data = {
    "sheet_name": "State Form",
    "print_area": "B1:Q107",
    "columns": {
        key: {
            "width": dim.width,
            "hidden": dim.hidden,
        }
        for key, dim in ws.column_dimensions.items()
    },
    "rows": {
        str(key): {
            "height": dim.height,
            "hidden": dim.hidden,
        }
        for key, dim in ws.row_dimensions.items()
    },
    "merged_cells": [
        str(r) for r in ws.merged_cells.ranges
        if r.min_col <= 17 and r.min_row <= 107
    ],
    "cells": {},
}

# Include A for source-only helper cells and R:AF for any
# historical supporting information, but these are not printed.
for row in ws.iter_rows(
    min_row=1,
    max_row=107,
    min_col=1,
    max_col=32,
):
    for cell in row:
        if cell.value is None and not cell.has_style:
            continue

        data["cells"][cell.coordinate] = {
            "value": cell.value,
            "style": style_dict(cell),
        }

DEST.write_text(json.dumps(data, indent=2, default=str))

print("Layout extracted:", DEST)
print("Cells:", len(data["cells"]))
print("Merged ranges:", len(data["merged_cells"]))
print("Historical workbook is NOT needed at runtime.")
